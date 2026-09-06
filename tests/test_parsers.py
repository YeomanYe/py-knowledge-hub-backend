"""解析器单元测试。"""
from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook

from app.parsers.docx_parser import parse_docx
from app.parsers.file_parser import FileParserService
from app.parsers.markdown_util import (
    clean_markdown,
    decode_upload_filename,
    escape_table_cell,
    get_extension,
    title_from_filename,
    to_markdown_table,
)
from app.parsers.plain_text_parser import parse_plain_text
from app.parsers.pptx_parser import parse_pptx
from app.parsers.xlsx_parser import parse_xlsx
from app.storage import RustfsService


class TestMarkdownUtil:
    def test_clean(self):
        assert clean_markdown("\r\na\r\n\r\n\r\n\r\nb  ") == "a\n\n\nb"

    def test_escape_cell(self):
        assert escape_table_cell("a|b\nc") == "a\\|b c"

    def test_table(self):
        out = to_markdown_table([["A", "B"], ["1", "2"]])
        assert "| A | B |" in out and "| --- | --- |" in out

    def test_extension_and_title(self):
        assert get_extension("a.PDF") == "pdf"
        assert title_from_filename("周报.docx") == "周报"

    def test_decode_filename(self):
        # Latin-1 误解码的 UTF-8 文件名可还原
        raw = "报告.docx".encode("utf-8").decode("latin-1")
        assert decode_upload_filename(raw) == "报告.docx"


class TestPlainText:
    def test_parse(self):
        assert parse_plain_text("你好\n世界".encode("utf-8")) == "你好\n世界"


class TestDocxParser:
    def test_parse(self):
        doc = Document()
        doc.add_heading("文档标题", level=0)
        doc.add_heading("第一章", level=1)
        doc.add_paragraph("这是一段正文，包含粗体。")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "姓名"
        table.cell(0, 1).text = "部门"
        table.cell(1, 0).text = "张三"
        table.cell(1, 1).text = "研发"
        buffer = io.BytesIO()
        doc.save(buffer)

        md = parse_docx(buffer.getvalue())
        assert "# 文档标题" in md
        assert "# 第一章" in md
        assert "这是一段正文" in md
        assert "| 姓名 | 部门 |" in md
        assert "| 张三 | 研发 |" in md


class TestPptxParser:
    def test_parse(self):
        slide = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
<p:cSld><p:spTree>
<p:sp><p:nvSpPr><p:cNvPr id="1" name="t"/><p:cNvSpPr/><p:nvPr><p:ph type="ctrTitle"/></p:nvPr></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>Hello 世界</a:t></a:r></a:p></p:txBody></p:sp>
<p:sp><p:nvSpPr><p:cNvPr id="2" name="b"/><p:cNvSpPr/><p:nvPr><p:ph type="body"/></p:nvPr></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>正文一</a:t></a:r></a:p><a:p><a:r><a:t>正文二</a:t></a:r></a:p></p:txBody></p:sp>
</p:spTree></p:cSld></p:sld>"""
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as zf:
            zf.writestr("ppt/slides/slide1.xml", slide)
            zf.writestr("ppt/slides/slide2.xml", slide.replace("正文一", "第二页内容"))
            zf.writestr("[Content_Types].xml", "<Types/>")
            zf.writestr("_rels/.rels", "<Relationships/>")

        md = parse_pptx(buffer.getvalue())
        assert "## 幻灯片 1" in md
        assert "### Hello 世界" in md
        assert "正文一" in md
        assert "## 幻灯片 2" in md
        assert "第二页内容" in md


class TestXlsxParser:
    def test_parse(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "人员"
        ws.append(["姓名", "数量"])
        ws.append(["张三", 3])
        ws.append(["李四", "=1+2"])
        buffer = io.BytesIO()
        wb.save(buffer)

        md = parse_xlsx(buffer.getvalue())
        assert "## 人员" in md
        assert "| 姓名 | 数量 |" in md
        assert "| 张三 | 3 |" in md


class TestFileParser:
    async def test_txt_dispatch(self):
        service = FileParserService(RustfsService())
        out = await service.parse("note.txt", "hello".encode())
        assert out == "hello"

    async def test_unsupported(self):
        service = FileParserService(RustfsService())
        with pytest.raises(Exception) as exc:
            await service.parse("a.exe", b"data")
        assert "不支持的文件格式" in str(exc.value)

    async def test_empty(self):
        service = FileParserService(RustfsService())
        with pytest.raises(Exception) as exc:
            await service.parse("a.txt", b"")
        assert "文件内容为空" in str(exc.value)
